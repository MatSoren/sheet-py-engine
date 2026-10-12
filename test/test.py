"""Differential check against Excel's own cached results.

Excel stores the last value it computed for every formula cell. Loading the same file
twice — once with the formulas, once with those cached values — gives a reference to
check the engine against, cell by cell.

    python -m test.test examples/example.xlsx
"""

import sys

from openpyxl import load_workbook

from translator import translator


def check_workbook(path, tolerance=6):
    """Returns the list of cells where the engine disagrees with Excel."""
    manager = translator.translate(load_workbook(path, data_only=False))
    expected_wb = load_workbook(path, data_only=True)

    mismatches = []
    for sheet_name in expected_wb.sheetnames:
        engine_sheet = manager.get_sheet(translator.normalize_sheet_name(sheet_name))
        if engine_sheet is None:
            continue
        for row in expected_wb[sheet_name].iter_rows():
            for cell in row:
                if cell.value is None or cell.value == "#REF!":
                    continue
                engine_cell = engine_sheet.get_cell(cell.coordinate)
                if engine_cell is None:
                    continue
                expected, actual = cell.value, engine_cell.calculated_value
                if isinstance(expected, (int, float)) and isinstance(actual, (int, float)):
                    failed = round(expected, tolerance) != round(actual, tolerance)
                else:
                    failed = expected != actual
                if failed:
                    mismatches.append((sheet_name, cell.coordinate, expected, actual))
    return mismatches


def main(argv=None):
    argv = argv if argv is not None else sys.argv[1:]
    if not argv:
        sys.exit(__doc__)

    failed = False
    for path in argv:
        mismatches = check_workbook(path)
        for sheet, coord, expected, actual in mismatches:
            print(f"MISMATCH {sheet}!{coord}: expected {expected!r}, got {actual!r}")
        print(f"{path}: {len(mismatches)} mismatch(es)")
        failed = failed or bool(mismatches)
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
