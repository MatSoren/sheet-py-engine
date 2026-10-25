from Sheet import SheetsManager, Sheet
from SheetCell import SheetCell, evaluate_in_dependency_order
from operations.operation_factory import OperationFactory, OperationTypes
from translator.AstProcessor import OperationProcessor, encode_error, encode_ref, normalize_sheet_name
from operations.operations import ExcelError
from openpyxl.formula.tokenizer import Tokenizer, Token
import ast
import logging
import re
import sys

sheet_manager = SheetsManager()
astProcessor = OperationProcessor()

# Excel infix operators that are spelled differently in Python. '&' becomes '|' on purpose:
# '|' sits between comparisons and '+' in Python's precedence, same as '&' in Excel.
infix_operator_map = {"=": "==", "<>": "!=", "^": "**", "&": "|"}


def formula_to_python(formula):
    """Excel formula ('=...') -> Python expression that ast.parse accepts.

    Goes token by token (openpyxl's tokenizer), so strings, references and operators are
    never confused with each other. OperationProcessor turns the parsed tree into operations.
    """
    out = []
    open_starts = []  # index in out where each open paren / function call started
    last_operand_start = 0  # index in out where the last complete operand started
    previous = None  # last non-whitespace token
    # Excel binds unary minus tighter than '^' and '^' is left-associative; Python does the
    # opposite, so both get explicit parentheses: -2^2 -> (-2)**2, 2^3^2 -> (2**3)**2
    prefix_starts = []  # (depth, index in out) of each '(-' waiting for its operand
    power_chain_start = {}  # depth -> index in out where the running '^' chain started
    items = Tokenizer(formula).items

    for index, token in enumerate(items):
        kind, subtype, value = token.type, token.subtype, token.value

        if kind == Token.WSPACE:
            following = next((t for t in items[index + 1:] if t.type != Token.WSPACE), None)
            if (
                previous is not None
                and following is not None
                and (previous.type == Token.OPERAND or previous.subtype == Token.CLOSE)
                and (following.type == Token.OPERAND or following.subtype == Token.OPEN)
            ):
                raise ValueError(f"intersection operator (space between ranges) not supported: {formula}")
            continue
        if kind == Token.ARRAY:
            raise ValueError(f"array constants ({{...}}) not supported: {formula}")

        # empty argument: IF(A1,,1) / IF(A1,1,)
        if previous is not None and (
            (kind == Token.SEP and (previous.type == Token.SEP or (previous.type == Token.FUNC and previous.subtype == Token.OPEN)))
            or (kind == Token.FUNC and subtype == Token.CLOSE and previous.type == Token.SEP)
        ):
            out.append("None")

        depth = len(open_starts)
        if kind == Token.OPERAND:
            last_operand_start = len(out)
            out.append(translate_operand(subtype, value))
        elif subtype == Token.OPEN:  # FUNC or PAREN
            open_starts.append(len(out))
            power_chain_start[depth + 1] = None
            out.append(value[:-1].upper() + "(" if kind == Token.FUNC else value)
        elif subtype == Token.CLOSE:
            last_operand_start = open_starts.pop()
            out.append(value)
        elif kind == Token.OP_PRE:
            prefix_starts.append((depth, len(out)))
            out.append("(" + value)
        elif kind == Token.OP_IN:
            if value == "^":
                chain_start = power_chain_start.get(depth)
                if chain_start is None:
                    power_chain_start[depth] = last_operand_start
                else:  # a^b^c: close (a**b) before adding **c
                    out.insert(chain_start, "(")
                    out.append(")")
                    last_operand_start = chain_start
            else:
                power_chain_start[depth] = None
            out.append(infix_operator_map.get(value, value))
        elif kind == Token.OP_POST:  # '%' binds to the operand right before it: 1/50% -> 1/(50/100)
            out.insert(last_operand_start, "(")
            out.append("/100)")
        else:  # SEP (',')
            power_chain_start[depth] = None
            out.append(value)

        # an operand just completed: close the '(-' prefixes waiting for it
        if kind == Token.OPERAND or subtype == Token.CLOSE:
            while prefix_starts and prefix_starts[-1][0] == len(open_starts):
                last_operand_start = prefix_starts.pop()[1]
                out.append(")")
        previous = token

    return "".join(out)


def translate_operand(subtype, value):
    if subtype == Token.TEXT:
        return repr(value[1:-1].replace('""', '"'))
    if subtype == Token.LOGICAL:
        return "True" if value.upper() == "TRUE" else "False"
    if subtype == Token.ERROR:
        return encode_error(value)
    if subtype == Token.NUMBER:
        return value
    # RANGE: A1, $A$1, A1:B2, Hoja!A1, 'Mi Hoja'!A1:B2
    if "#REF!" in value:
        return encode_error("#REF!")
    start, colon, end = value.partition(":")  # sheet names can't contain ':'
    if "!" in end:
        if "!" not in start:
            raise ValueError(f"3D references (Sheet1:Sheet3!A1) not supported: {value}")
        value = start + colon + end.rpartition("!")[2]  # Hoja!A1:Hoja!B2 -> Hoja!A1:B2
    sheet, _, ref = value.rpartition("!")
    if sheet.startswith("'"):
        sheet = sheet[1:-1].replace("''", "'")
    return encode_ref(sheet or None, ref.replace("$", ""))


def translate(workbook):
    creation_switch = {
        "n": generate_value_cell,
        "s": generate_value_cell,
        "f": generate_formula_cell,
        "b": generate_value_cell,
        "d": generate_value_cell,
        "e": generate_value_cell,
    }
    normalized_names = [normalize_sheet_name(name) for name in workbook.sheetnames]
    if len(set(normalized_names)) != len(normalized_names):
        raise ValueError(f"sheet names collide once normalized: {workbook.sheetnames}")
    for sheet_name in normalized_names:
        sheet_manager.instantiate_sheet(sheet_name)
    for raw_name in workbook.sheetnames:
        sheet = workbook[raw_name]
        sheet_name = normalize_sheet_name(raw_name)

        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                creation_switch.get(cell.data_type, None)(cell, sheet_name)
    evaluate_in_dependency_order(
        [cell for name in normalized_names for cell in sheet_manager.get_sheet(name).elements.values()]
    )
    return sheet_manager


def generate_value_cell(cell, sheet_name):
    sheet = sheet_manager.get_sheet(sheet_name)
    sheet_cell = sheet.get_cell(cell.coordinate)
    value = ExcelError(cell.value) if cell.data_type == "e" else cell.value
    if sheet_cell == None:
        sheet_cell = SheetCell(OperationTypes.Value, value)
        sheet_cell.set_coordinate(cell.coordinate)
        sheet_cell.set_sheet(sheet_name)

    else:
        sheet_cell.operation = OperationFactory.get_instance(OperationTypes.Value, value)
    sheet.set_cell(cell.coordinate, sheet_cell)


formule_generators = {}


def generate_countif():
    return OperationTypes.Value, 0.000101


def generate_if():
    return OperationTypes.Value, 0.000101


def generate_mround():
    return OperationTypes.Value, 0.000101


def generate_rounddown():
    return OperationTypes.Value, 0.000101


def generate_roundup():
    return OperationTypes.Value, 0.000101


def generate_sum():
    return OperationTypes.Value, 0.000101


def generate_today():
    return OperationTypes.Value, 0.000101


def generate_iferror():
    return OperationTypes.Value, 0.000101


def generate_text():
    return OperationTypes.Value, 0.000101


formule_generators = {
    "COUNTIF": generate_countif,
    "IF": generate_if,
    "MROUND": generate_mround,
    "ROUNDDOWN": generate_rounddown,
    "ROUNDUP": generate_roundup,
    "SUM": generate_sum,
    "TODAY": generate_today,
    "IFERROR": generate_iferror,
    "TEXT": generate_text,
}


def generate_formula_cell(cell, sheet_name):
    current_sheet: Sheet = sheet_manager.get_sheet(sheet_name)

    current_cell = current_sheet.get_cell(cell.coordinate)

    if current_cell is None:
        current_cell = SheetCell()

        current_sheet.set_cell(cell.coordinate, current_cell)
    astProcessor.set_current_cell(current_cell)
    astProcessor.set_current_sheet(current_sheet)

    # a formula the engine can't build is an error in its own cell, not a failed load
    try:
        tree = ast.parse(formula_to_python(str(cell.value)))
        operation = astProcessor.visit(tree)
    except (SyntaxError, ValueError, TypeError) as error:
        logging.warning("cannot evaluate %s!%s %s: %s", sheet_name, cell.coordinate, cell.value, error)
        operation = OperationFactory.get_instance(OperationTypes.Value, ExcelError("#VALUE!"))

    current_cell.operation = operation  # computed by evaluate_in_dependency_order()


def have_operators(expression):
    pattern = r"[+\-*/()<> =]"
    return bool(re.search(pattern, expression))


def generate_expression_operation(expression: str, sheet_name):
    sheet_referenced_name = sheet_name
    cell_name = expression

    if expression.find("!") != -1:
        sheet_referenced_name, cell_name = expression.split("!")

    sheet = sheet_manager.get_sheet(sheet_referenced_name)
    cell_referenced = sheet.get_cell(cell_name)
    if cell_referenced is None:
        cell_referenced = SheetCell()
        sheet.set_cell(cell_name, cell_referenced)
    return OperationTypes.Reference, cell_referenced
