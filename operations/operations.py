import logging
from abc import ABC, abstractmethod

from datetime import date, datetime, timedelta

import openpyxl.utils.cell as utils

import re

import math

from decimal import Decimal, ROUND_HALF_UP

from collections.abc import Iterable

"""
Operaciones:
- COUNTIF
- IF *
- MROUND *
- ROUNDDOWN *
- ROUNDUP *
- SUM
- TODAY*
- iferror
- text *
- +
- -
- *
- /
"""


class Operation(ABC):
    def __str__(self):
        return "Operation"

    @abstractmethod
    def operate(self):
        pass


class NoneOperation(Operation):
    def __str__(self):
        return "NoneOperation"

    def __init(self):
        pass

    def operate(self):
        return None


class ExcelError(Exception):
    """An Excel error value (#N/A, #DIV/0!, ...).

    Cells hold it as their value; operations raise it, so it propagates through any formula
    until IFERROR or the cell itself catches it, like in Excel. It compares equal to its
    code: ExcelError("#N/A") == "#N/A".
    """

    def __eq__(self, other):
        if isinstance(other, (str, ExcelError)):
            return str(self) == str(other)
        return NotImplemented

    def __hash__(self):
        return hash(str(self))

    @classmethod
    def from_exception(cls, error):
        """A formula failure -> the error Excel would show for it."""
        if isinstance(error, ExcelError):
            return error
        if isinstance(error, ZeroDivisionError):
            return cls("#DIV/0!")
        if isinstance(error, OverflowError):
            return cls("#NUM!")
        return cls("#VALUE!")


# what evaluating a formula can fail with on spreadsheet data; anything else is a bug
FORMULA_FAILURES = (ExcelError, ArithmeticError, ValueError, TypeError)


def raise_if_error(value):
    if isinstance(value, ExcelError):
        raise ExcelError(str(value))  # a fresh one: re-raising would keep growing its traceback
    return value


class Value(Operation):
    def __str__(self):
        return "Value"

    def __init__(self, value):
        self.value = value

    def operate(self):
        return raise_if_error(self.value)


class Reference(Operation):
    def __str__(self):
        return "Reference"

    def __init__(self, cell):
        self.cell = cell

    def operate(self):
        return raise_if_error(self.cell.calculated_value)


class Range(Operation):
    def __str__(self):
        return "Range"

    def __init__(self, sheet_range, current_sheet):
        initial, final = sheet_range.split(":")
        initial_row, initial_col = utils.coordinate_to_tuple(initial)
        final_row, final_col = utils.coordinate_to_tuple(final)
        self.width = final_col - initial_col + 1

        self.cell_range = [
            current_sheet.get_cell(utils.get_column_letter(col) + str(row))
            for row in range(initial_row, final_row + 1)
            for col in range(initial_col, final_col + 1)
        ]

    def operate(self):
        return self.cell_range


class Today(Operation):
    def __str__(self):
        return "Today"

    def __init__(self):
        pass

    def operate(self):
        return datetime.today().replace(hour=0, minute=0, second=0, microsecond=0)


def excel_comparable(left, right):
    """Two values -> two keys Python compares the way Excel does.

    numbers < text < TRUE/FALSE, text is case-insensitive, dates are serial numbers, and a
    blank cell counts as 0, "" or FALSE depending on the other side.
    """
    if left is None:
        left = _blank_like(right)
    if right is None:
        right = _blank_like(left)
    return _comparison_key(left), _comparison_key(right)


def _blank_like(other):
    if isinstance(other, bool):
        return False
    if isinstance(other, str):
        return ""
    return 0


def _comparison_key(value):
    if isinstance(value, bool):
        return (2, value)
    if isinstance(value, str):
        return (1, value.lower())
    if isinstance(value, datetime):
        return (0, (value - datetime(1899, 12, 30)).total_seconds() / 86400)
    if isinstance(value, date):
        return (0, (value - date(1899, 12, 30)).days)
    return (0, value)


class EqualCondition(Operation):
    def __str__(self):
        return "EqualCondition"

    def __init__(self, operand_1, operand_2):
        self.operand_1 = operand_1
        self.operand_2 = operand_2

    def operate(self):
        left, right = excel_comparable(self.operand_1.operate(), self.operand_2.operate())
        return left == right


class DiferentCondition(Operation):
    def __str__(self):
        return "DiferentCondition"

    def __init__(self, operand_1, operand_2):
        self.operand_1 = operand_1
        self.operand_2 = operand_2

    def operate(self):
        left, right = excel_comparable(self.operand_1.operate(), self.operand_2.operate())
        return left != right


class GratherCondition(Operation):
    def __str__(self):
        return "GratherCondition"

    def __init__(self, operand_1, operand_2):
        self.operand_1 = operand_1
        self.operand_2 = operand_2

    def operate(self):
        left, right = excel_comparable(self.operand_1.operate(), self.operand_2.operate())
        return left > right


class GratherOrEqualCondition(Operation):
    def __str__(self):
        return "GratherOrEqualCondition"

    def __init__(self, operand_1, operand_2):
        self.operand_1 = operand_1
        self.operand_2 = operand_2

    def operate(self):
        left, right = excel_comparable(self.operand_1.operate(), self.operand_2.operate())
        return left >= right


class LessCondition(Operation):
    def __str__(self):
        return "LessCondition"

    def __init__(self, operand_1, operand_2):
        self.operand_1 = operand_1
        self.operand_2 = operand_2

    def operate(self):
        left, right = excel_comparable(self.operand_1.operate(), self.operand_2.operate())
        return left < right


class LessOrEqualCondition(Operation):
    def __str__(self):
        return "LessOrEqualCondition"

    def __init__(self, operand_1, operand_2):
        self.operand_1 = operand_1
        self.operand_2 = operand_2

    def operate(self):
        left, right = excel_comparable(self.operand_1.operate(), self.operand_2.operate())
        return left <= right


class BlankCondition(Operation):
    def __str__(self):
        return "BlankCondition"

    def __init__(self, cell):
        self.cell = cell

    def operate(self):
        if self.cell.operate() is None:
            return True
        return False


class AndCondition(Operation):
    def __str__(self):
        return "AndCondition"

    def __init__(self, *args, **kargs):
        self.operands = list(args)

    def operate(self):
        acumulated_bool = True
        for op in self.operands:
            operated = op.operate()
            if operated is None:
                return None

            acumulated_bool = acumulated_bool and operated
            if not acumulated_bool:
                return acumulated_bool

        return acumulated_bool


class OrCondition(Operation):
    def __str__(self):
        return "OrCondition"

    def __init__(self, *args, **kargs):
        self.operands = list(args)

    def operate(self):

        acumulated_bool = False
        for op in self.operands:
            operated = op.operate()
            if operated is None:
                return None

            acumulated_bool = acumulated_bool or operated

        return acumulated_bool


class NotCondition(Operation):
    def __str__(self):
        return "NotCondition"

    def __init__(self, operand):
        self.operand = operand

    def operate(self):
        return not self.operand.operate()


class If(Operation):
    def __str__(self):
        return "If"

    def __init__(self, _condition, _then, _else=Value(False)):
        self._condition = _condition
        self._then = _then
        self._else = _else

    def operate(self):
        if self._condition.operate():
            return self._then.operate()
        return self._else.operate()


class CountIf(Operation):
    def __str__(self):
        return "CountIf"

    _condition_operation = {
        "<=": lambda left, right: left <= right,
        ">=": lambda left, right: left >= right,
        "<>": lambda left, right: left != right,
        "=": lambda left, right: left == right,
        "<": lambda left, right: left < right,
        ">": lambda left, right: left > right,
    }

    def __init__(self, cell_range, condition):
        self.range = cell_range
        self.condition = condition

    def operate(self):
        # evaluated here, not in __init__: the criterion may live in a cell loaded later
        matches = self._matcher(*self._split_condition(self.condition.operate()))
        # COUNTIF(A1, ...) gets a Reference; read the cell, an error in it just doesn't match
        cells = self.range.operate() if isinstance(self.range, Range) else [self.range.cell]
        return sum(1 for cell in cells if matches(cell.calculated_value))

    def _split_condition(self, criteria):
        """'>=10' -> ('>=', 10.0), 'abc' -> ('=', 'abc'), 5 -> ('=', 5); empty counts as 0."""
        if criteria is None:
            return "=", 0
        if not isinstance(criteria, str):
            return "=", criteria
        operator, target = re.match(r"(<=|>=|<>|=|<|>)?(.*)", criteria, re.DOTALL).groups()
        try:
            target = float(target)
        except ValueError:
            pass
        return operator or "=", target

    def _matcher(self, operator, target):
        operation = self._condition_operation[operator]
        equality = operator in ("=", "<>")
        if isinstance(target, str):
            pattern = wildcard_pattern(target)

            def matches(value):
                if value is None or isinstance(value, ExcelError):
                    # a blank only equals "" and an error only its code: "*" counts neither
                    code = "" if value is None else str(value)
                    return equality and (code.lower() == target.lower()) == (operator == "=")
                if not isinstance(value, str):
                    return operator == "<>"
                if equality:
                    return bool(pattern.fullmatch(value)) == (operator == "=")
                return operation(value.lower(), target.lower())

            return matches

        def matches(value):
            if isinstance(value, str) and equality:
                try:  # a number stored as text equals a numeric criterion, like in Excel
                    value = float(value)
                except ValueError:
                    pass
            if (
                value is None
                or isinstance(value, (str, ExcelError))
                or isinstance(value, bool) != isinstance(target, bool)
            ):
                return operator == "<>"
            return operation(value, target)

        return matches


def wildcard_pattern(text):
    """COUNTIF text criteria: * any run, ? one character, ~ escapes the next; any case."""
    pattern = "".join(
        ".*" if part == "*" else "." if part == "?" else re.escape(part[-1])
        for part in re.findall(r"~.|.", text, re.DOTALL)
    )
    return re.compile(pattern, re.IGNORECASE | re.DOTALL)


class MRound(Operation):
    def __str__(self):
        return "MRound"

    def __init__(self, number, multiple):
        self.number = number
        self.multiple = multiple

    def operate(self):
        number = self.number.operate()
        mutiple = self.multiple.operate()
        if mutiple is None or number is None:
            return None
        if mutiple == 0:
            return 0
        # Excel rounds halves away from zero; Python's round() goes to the even neighbour.
        # ponytail: 12 decimals absorb float noise (0.15/0.1 = 1.4999999999999998)
        quotient = round(number / mutiple, 12)
        return math.copysign(math.floor(abs(quotient) + 0.5), quotient) * mutiple


def round_magnitude(number, num_digits, rounding):
    """ROUNDUP/ROUNDDOWN: apply floor/ceil to the absolute value, so the sign doesn't flip
    the direction. 12 decimals absorb float noise first (1.1*3 = 3.3000000000000003)."""
    factor = 10**num_digits
    result = math.copysign(rounding(round(abs(number) * factor, 12)) / factor, number)
    return result or 0.0  # no -0.0


class RoundDown(Operation):
    def __str__(self):
        return "RoundDown"

    def __init__(self, number, num_digits=Value(0)):
        self.number = number
        self.num_digits = num_digits

    def operate(self):
        number = self.number.operate()
        if self.num_digits.operate() is None or number is None:
            return None
        return round_magnitude(number, self.num_digits.operate(), math.floor)  # toward zero


class RoundUp(Operation):
    def __str__(self):
        return "RoundUp"

    def __init__(self, number, num_digits=Value(0)):
        self.number = number
        self.num_digits = num_digits

    def operate(self):
        number = self.number.operate()
        if self.num_digits.operate() is None or number is None:
            return None
        return round_magnitude(number, self.num_digits.operate(), math.ceil)  # away from zero


class BinaryOperation(Operation, ABC):
    def __str__(self):
        return "BinaryAdd"

    def __init__(self, left, right):
        self.left = left
        self.right = right

    def operands(self):
        """Both sides evaluated; a blank cell counts as 0 in arithmetic, like in Excel."""
        return [0 if value is None else value for value in (self.left.operate(), self.right.operate())]

    @staticmethod
    def as_number(value):
        return float(value) if isinstance(value, str) else value


class BinaryAdd(BinaryOperation):
    def operate(self):
        left, right = self.operands()
        if isinstance(left, datetime) or isinstance(right, datetime):
            return self.add_dates(left, right)
        try:
            left = float(left)
            right = float(right)
        except:
            left = str(left)
            right = str(right)
        return left + right

    def add_dates(self, left, right):
        number = ""
        date = ""

        if isinstance(left, (int, float)):
            number = left
            date = right

        if isinstance(right, (int, float)):
            number = right
            date = left

        return date + timedelta(days=number)


class BinarySub(BinaryOperation):
    def operate(self):
        left, right = self.operands()
        return left - right


class BinaryMul(BinaryOperation):
    def operate(self):
        left, right = map(self.as_number, self.operands())
        return left * right


class BinaryDiv(BinaryOperation):
    def operate(self):
        left, right = map(self.as_number, self.operands())
        return left / right  # ZeroDivisionError becomes #DIV/0! in the cell


class BinaryPow(BinaryOperation):
    def operate(self):
        left, right = map(self.as_number, self.operands())
        return left**right


def to_excel_text(value):
    """How Excel shows a value inside a concatenation: 1.0 -> '1', True -> 'TRUE'."""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return format(value, ".15g")
    return str(value)


class Concat(BinaryOperation):
    def operate(self):
        return to_excel_text(self.left.operate()) + to_excel_text(self.right.operate())


class Sum(Operation):
    def __str__(self):
        return "Sum"

    def __init__(self, cell_range):
        self.range = cell_range

    def operate(self):
        sum_calc = 0
        _range = self.range.operate()
        if not isinstance(_range, Iterable):
            return _range
        for cell in _range:
            if cell is None:
                continue
            cell_value = raise_if_error(cell.calculated_value)
            if isinstance(cell_value, (int, float, complex)):
                sum_calc += cell_value
        return sum_calc


class Hlookup(Operation):
    def __str__(self):
        return "Hlookup"

    def __init__(self, lookup_value, cell_range, row_index, approximated=Value(True)):
        self.lookup_value = lookup_value
        self.cell_range = cell_range
        self.row_index = row_index
        self.approximated = approximated

    def operate(self):
        lookup = self.lookup_value.operate()
        row_index = self.row_index.operate()
        if lookup is None or row_index is None:
            return None
        cells = self.cell_range.operate()
        width = self.cell_range.width
        first_row = [cell.calculated_value for cell in cells[:width]]
        column = self._find_column(first_row, lookup, self.approximated.operate())
        index = (int(row_index) - 1) * width + (column or 0)
        if column is None:
            raise ExcelError("#N/A")
        if not 0 <= index < len(cells):
            raise ExcelError("#REF!")
        return raise_if_error(cells[index].calculated_value)

    @staticmethod
    def _find_column(first_row, lookup, approximated):
        """Exact: first match. Approximate: last value <= lookup, assuming an ascending row.

        Like Excel, values of the other kind (text vs number) are skipped, text is
        case-insensitive.
        """
        is_text = isinstance(lookup, str)
        key = lookup.lower() if is_text else lookup
        found = None
        for column, value in enumerate(first_row):
            if value is None or isinstance(value, ExcelError) or isinstance(value, str) != is_text:
                continue
            value = value.lower() if is_text else value
            if value == key:
                return column
            if approximated:
                if value > key:
                    break
                found = column
        return found if approximated else None


# one token of an Excel format code: literals ("..", \x, [..]), date codes, digit placeholders
FORMAT_TOKEN = re.compile(
    r'"[^"]*"|\\.|\[[^\]]*\]|am/pm|a/p|yyyy|yy|mmmm|mmm|mm|m|dddd|ddd|dd|d|hh|h|ss|s|[0#?]|.',
    re.IGNORECASE,
)
DATE_CODES = {"yyyy", "yy", "mmmm", "mmm", "mm", "m", "dddd", "ddd", "dd", "d", "hh", "h", "ss", "s", "am/pm", "a/p"}


def format_literal(token):
    if token.startswith('"'):
        return token[1:-1]
    if token.startswith("\\"):
        return token[1:]
    if token.startswith("["):
        return ""  # [Red], [$-409], ...
    return token


class Text(Operation):
    def __str__(self):
        return "Text"

    def __init__(self, element, _format):
        self.element = element
        self.format = _format

    def operate(self):
        # evaluated here, not in __init__, so TEXT follows changes to its source cell
        element = self.element.operate()
        _format = self.format.operate()
        if isinstance(element, bool):
            return "TRUE" if element else "FALSE"
        if isinstance(element, str):
            return element
        if element is None:
            element = 0
        sections = [[]]
        for token in FORMAT_TOKEN.findall(_format):
            if token == ";":
                sections.append([])
            else:
                sections[-1].append(token)
        is_number = not isinstance(element, (date, datetime))
        if is_number and element < 0 and len(sections) > 1:
            tokens, element, sign = sections[1], -element, ""  # the section writes its own sign
        elif is_number and element == 0 and len(sections) > 2:
            tokens, sign = sections[2], ""
        else:
            tokens, sign = sections[0], "-" if is_number and element < 0 else ""
        if any(token.lower() in DATE_CODES for token in tokens):
            if is_number:
                element = datetime(1899, 12, 30) + timedelta(days=element)  # serial date
            return self._date_text(element, tokens)
        if not is_number:  # a date shown with a number format: its serial number
            element = _comparison_key(element)[1]
        return self._number_text(abs(element), tokens, sign)

    @staticmethod
    def _date_text(moment, tokens):
        if not isinstance(moment, datetime):
            moment = datetime(moment.year, moment.month, moment.day)
        twelve_hour = any(token.lower() in ("am/pm", "a/p") for token in tokens)
        codes = [(index, token.lower()) for index, token in enumerate(tokens) if token.lower() in DATE_CODES]
        out = [format_literal(token) for token in tokens]
        for position, (index, code) in enumerate(codes):
            if code in ("m", "mm"):  # minutes right after an hour or right before seconds
                before = codes[position - 1][1] if position > 0 else ""
                after = codes[position + 1][1] if position + 1 < len(codes) else ""
                if before in ("h", "hh") or after in ("s", "ss"):
                    code = "minute" + code
            hour = (moment.hour % 12 or 12) if twelve_hour else moment.hour
            out[index] = {
                "yyyy": f"{moment.year:04d}", "yy": f"{moment.year % 100:02d}",
                "mmmm": moment.strftime("%B"), "mmm": moment.strftime("%b"),
                "mm": f"{moment.month:02d}", "m": str(moment.month),
                "minutemm": f"{moment.minute:02d}", "minutem": str(moment.minute),
                "dddd": moment.strftime("%A"), "ddd": moment.strftime("%a"),
                "dd": f"{moment.day:02d}", "d": str(moment.day),
                "hh": f"{hour:02d}", "h": str(hour),
                "ss": f"{moment.second:02d}", "s": str(moment.second),
                "am/pm": "AM" if moment.hour < 12 else "PM", "a/p": "A" if moment.hour < 12 else "P",
            }[code]
        return "".join(out)

    @staticmethod
    def _number_text(value, tokens, sign):
        """Digit placeholders (0 # ?), '.', ',' (grouping, or /1000 when trailing) and '%'
        (x100); everything else is copied as a literal.

        ponytail: the digits go where the first placeholder is, so literals between
        placeholders (000-0000) and '?' space padding are not supported.
        """
        if [token.lower() for token in tokens] == list("general"):
            return sign + to_excel_text(value)
        body, out, number_at = "", [], None
        for token in tokens:
            # '.' may start the number (".00"), ',' only continues it
            if token in ("0", "#", "?", ".") or (token == "," and number_at is not None):
                if number_at is None:
                    number_at = len(out)
                body += token
            else:
                out.append(format_literal(token))
        if number_at is None:
            return "".join(out)
        integer_spec, _, decimal_spec = body.partition(".")
        # commas closing a part scale by 1000 each: #,##0, / 0.0,,
        scale_commas = sum(len(spec) - len(spec.rstrip(",")) for spec in (integer_spec, decimal_spec))
        integer_spec, decimal_spec = integer_spec.rstrip(","), decimal_spec.rstrip(",")
        number = Decimal(repr(value)) * 100 ** tokens.count("%") / 1000**scale_commas
        decimals = len(decimal_spec)
        rounded = number.quantize(Decimal(1).scaleb(-decimals), ROUND_HALF_UP)
        integer_digits, _, decimal_digits = f"{rounded:f}".partition(".")
        minimum_digits = integer_spec.count("0")
        integer_digits = integer_digits.lstrip("0").zfill(minimum_digits)
        if "," in integer_spec:
            integer_digits = re.sub(r"(?<=\d)(?=(\d{3})+$)", ",", integer_digits)
        decimal_digits = decimal_digits.rstrip("0").ljust(decimal_spec.count("0"), "0")
        text = integer_digits + ("." + decimal_digits if "." in body else "")
        out.insert(number_at, sign + text)
        return "".join(out)


class IfError(Operation):
    def __str__(self):
        return "IfError"

    def __init__(self, _then, _else):
        self._then = _then
        self._else = _else

    def operate(self):
        try:
            return self._then.operate()
        except FORMULA_FAILURES:
            return self._else.operate()
