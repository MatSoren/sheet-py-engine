import openpyxl.utils.cell as utils

from SheetCell import SheetCell


class SheetsManager:
    _sheet_dict = {}
    manager = None

    def __init__(self):
        pass

    def instantiate_sheet(self, sheet_name):
        sheet = Sheet(sheet_name)
        SheetsManager._sheet_dict[sheet_name] = sheet
        return sheet

    def set_sheet(self, sheet_name, sheet_data):
        SheetsManager._sheet_dict[sheet_name] = sheet_data

    def get_sheet(self, sheet_name):
        return SheetsManager._sheet_dict.get(sheet_name, None)


class Sheet:
    def __init__(self, name=""):
        self.sheet_name = name
        self.elements = {}
        self.max_row = self.max_column = 0  # the last row / column holding a cell
        self.open_ranges = []  # whole-column / whole-row ranges (A:A, 1:1) over this sheet

    def get_cell(self, cell_id):
        return self.elements.get(cell_id, None)

    def cell_at(self, cell_id):
        """The cell at cell_id, registering a blank one if there is none yet."""
        cell = self.elements.get(cell_id)
        if cell is None:
            cell = SheetCell()
            self.set_cell(cell_id, cell)
        return cell

    def set_cell(self, cell_id, cell):
        self.elements[cell_id] = cell
        cell.set_coordinate(cell_id)
        cell.set_sheet(self.sheet_name)
        row, column = utils.coordinate_to_tuple(cell_id)
        grew = row > self.max_row or column > self.max_column
        self.max_row, self.max_column = max(self.max_row, row), max(self.max_column, column)
        for open_range in self.open_ranges:
            if open_range.contains(row, column):
                cell.attach(open_range.observer)  # a cell new to A:A notifies the formula over it
                open_range.cell_range = None
            elif grew:
                open_range.cell_range = None  # its open side now ends further out

    def __str__(self):
        return "\n".join(
            f"{element_id}: {element}" for element_id, element in self.elements.items()
        )
