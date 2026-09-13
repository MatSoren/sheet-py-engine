"""CLI for the spreadsheet engine.

Loads an .xlsx file, builds the reactive cell graph from its formulas, and lets you
read cells or overwrite values and watch the recalculation propagate.
"""

import argparse
import logging
import sys

from openpyxl import load_workbook

from operations.operations import Value
from translator import translator


def parse_target(target):
    """'SHEET!A1' -> ('SHEET', 'A1')"""
    if "!" not in target:
        raise argparse.ArgumentTypeError(
            f"expected SHEET!CELL (e.g. Budget!F5), got {target!r}"
        )
    sheet, _, cell = target.partition("!")
    return sheet, cell


def resolve(manager, sheet_name, coordinate):
    sheet = manager.get_sheet(sheet_name)
    if sheet is None:
        sys.exit(f"no such sheet: {sheet_name!r}")
    cell = sheet.get_cell(coordinate)
    if cell is None:
        sys.exit(f"no such cell: {sheet_name}!{coordinate}")
    return cell


def coerce(raw):
    """Numbers stay numbers, everything else is a string."""
    try:
        return int(raw)
    except ValueError:
        pass
    try:
        return float(raw)
    except ValueError:
        return raw


def main(argv=None):
    parser = argparse.ArgumentParser(
        prog="sheet-engine", description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("xlsx", help="path to the .xlsx file")
    parser.add_argument(
        "cells", nargs="*", metavar="SHEET!CELL",
        help="cells to print after loading (e.g. Budget!F5)",
    )
    parser.add_argument(
        "--set", action="append", default=[], metavar="SHEET!CELL=VALUE",
        help="overwrite a cell before reading; repeatable. Recalculation propagates.",
    )
    parser.add_argument(
        "--recursion-limit", type=int, default=1000,
        help="Python recursion limit; deep formula chains need a high value (default: 1000)",
    )
    parser.add_argument("-v", "--verbose", action="store_true", help="log the propagation")
    args = parser.parse_args(argv)

    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.WARNING,
        format="%(levelname)s %(message)s",
    )
    sys.setrecursionlimit(args.recursion_limit)

    workbook = load_workbook(args.xlsx, data_only=False)
    manager = translator.translate(workbook)

    for assignment in args.set:
        target, _, raw = assignment.partition("=")
        if not raw:
            sys.exit(f"expected SHEET!CELL=VALUE, got {assignment!r}")
        sheet_name, coordinate = parse_target(target)
        cell = resolve(manager, sheet_name, coordinate)
        # set_operation triggers update() -> notify(), cascading to every dependent cell
        cell.set_operation(Value(coerce(raw)))
        print(f"set {sheet_name}!{coordinate} = {cell.calculated_value}")

    for target in args.cells:
        sheet_name, coordinate = parse_target(target)
        cell = resolve(manager, sheet_name, coordinate)
        print(f"{sheet_name}!{coordinate} = {cell.calculated_value}")

    if not args.cells and not args.set:
        print("sheets:", ", ".join(workbook.sheetnames))


if __name__ == "__main__":
    main()
